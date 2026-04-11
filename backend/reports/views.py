import openpyxl
from django.http import HttpResponse
from rest_framework import permissions
from rest_framework.views import APIView

from bookings.models import Booking
from crm.models import Customer
from core.utils import (
    get_user_profile,
    get_user_restaurant,
    auto_adjust_column_width,
    create_excel_response,
)

from core.permissions import CanViewCustomers, IsRestaurantStaff


class BookingsExcelReportView(APIView):
    permission_classes = [permissions.IsAuthenticated, IsRestaurantStaff]

    def get(self, request):
        restaurant = get_user_restaurant(request.user)
        if not restaurant:
            return HttpResponse("Unauthorized", status=403)

        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet(title="Bookings")
        headers = ["ID", "Date", "Time", "Guests", "Customer Name", "Phone", "Status", "Table"]
        ws.append(headers)
        
        bookings = (
            Booking.objects.filter(restaurant=restaurant)
            .select_related("user", "user__profile", "table")
            .order_by("-date", "-time")
        ).iterator(chunk_size=2000)
        
        for b in bookings:
            customer_name = b.user_name or (b.user.get_full_name().strip() if b.user else "") or (b.user.username if b.user else "") or "Guest"
            user_profile = get_user_profile(b.user) if b.user else None
            customer_phone = b.user_phone or (getattr(user_profile, "phone", "") if user_profile else "") or ""
            ws.append([
                b.id,
                b.date.strftime("%Y-%m-%d") if b.date else "",
                b.time.strftime("%H:%M") if b.time else "",
                b.guests,
                customer_name,
                customer_phone,
                b.status,
                str(b.table.number) if b.table else "",
            ])
            
        return create_excel_response(wb, "Bookings_Report.xlsx")


class CustomersExcelReportView(APIView):
    permission_classes = [permissions.IsAuthenticated, CanViewCustomers]

    def get(self, request):
        restaurant = get_user_restaurant(request.user)
        if not restaurant:
            return HttpResponse("Unauthorized", status=403)

        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet(title="Customers")
        headers = ["ID", "Name", "Phone", "Email", "Visits Count", "Total Spent", "Notes Count", "Last Note"]
        ws.append(headers)
        
        customers = Customer.objects.filter(restaurant=restaurant).prefetch_related("internal_notes").order_by("-visits_count").iterator(chunk_size=2000)
        
        for c in customers:
            notes = list(c.internal_notes.all())
            last_note = max(notes, key=lambda n: n.created_at).content if notes else ""
            ws.append([
                c.id,
                c.name,
                c.phone,
                c.email or "",
                c.visits_count,
                float(c.total_spent),
                len(notes),
                last_note,
            ])
            
        return create_excel_response(wb, "Customers_Report.xlsx")


class AnalyticsExcelReportView(APIView):
    permission_classes = [permissions.IsAuthenticated, IsRestaurantStaff]

    def get(self, request):
        restaurant = get_user_restaurant(request.user)
        if not restaurant:
            return HttpResponse("Unauthorized", status=403)

        wb = openpyxl.Workbook(write_only=True)
        ws = wb.create_sheet(title="Analytics Overview")
        headers = ["Metric", "Value"]
        ws.append(headers)
        total_bookings = Booking.objects.filter(restaurant=restaurant, status="completed").count()
        total_customers = Customer.objects.filter(restaurant=restaurant).count()
        avg_check = restaurant.average_price or 0
        total_revenue = total_bookings * avg_check
        ws.append(["Total Completed Bookings", total_bookings])
        ws.append(["Total Customers", total_customers])
        ws.append(["Average Check", avg_check])
        ws.append(["Estimated Revenue", total_revenue])
        return create_excel_response(wb, "Analytics_Report.xlsx")
